"""Report generation service — PDF and Excel exports."""

import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from models.prediction import PredictionRecord
from models.project import Project
from services.audit_service import AuditService


class ReportService:
    """Generates downloadable PDF and Excel reports."""

    @staticmethod
    def _get_prediction_or_project(
        db: Session, prediction_id: Optional[str] = None, project_id: Optional[str] = None,
    ) -> tuple[Optional[PredictionRecord], Optional[Project]]:
        prediction = None
        project = None
        if prediction_id:
            prediction = db.query(PredictionRecord).filter(PredictionRecord.id == prediction_id).first()
            if not prediction:
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Prediction not found.")
        if project_id:
            project = db.query(Project).filter(Project.id == project_id).first()
            if not project:
                raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Project not found.")
        return prediction, project

    @staticmethod
    def generate_excel(
        db: Session,
        user_id: str,
        prediction_id: Optional[str] = None,
        project_id: Optional[str] = None,
        ip_address: Optional[str] = None,
    ) -> bytes:
        """Generate an Excel report with prediction summary and risk breakdown."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            raise HTTPException(status_code=500, detail="openpyxl is required for Excel reports.")

        prediction, project = ReportService._get_prediction_or_project(db, prediction_id, project_id)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Risk Report"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

        ws["A1"] = "RiskVision AI — Risk Assessment Report"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Generated: {datetime.now(timezone.utc).isoformat()}"

        row = 4
        headers = ["Field", "Value"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill

        data_rows = []
        if prediction:
            data_rows = [
                ("Project ID", prediction.external_project_id),
                ("Project Name", prediction.project_name or "N/A"),
                ("Risk Level", prediction.risk_level),
                ("Risk Score", prediction.risk_score),
                ("Failure Probability", f"{prediction.failure_probability:.2%}"),
                ("Prediction Label", prediction.prediction_label),
                ("Confidence", f"{prediction.confidence_level:.2%}"),
                ("Model Version", prediction.model_version or "N/A"),
                ("Predicted At", prediction.predicted_at.isoformat()),
            ]
            for factor in (prediction.top_risk_factors or [])[:5]:
                if isinstance(factor, dict):
                    data_rows.append((
                        f"Risk Factor: {factor.get('display_name', '')}",
                        f"Impact: {factor.get('impact', 0):.4f} ({factor.get('direction', '')})",
                    ))
        elif project:
            data_rows = [
                ("Project ID", project.external_id),
                ("Project Name", project.name),
                ("Status", project.status),
                ("Latest Risk Level", project.latest_risk_level or "N/A"),
                ("Latest Risk Score", project.latest_risk_score or "N/A"),
                ("Budget", project.budget),
                ("Team Size", project.team_size),
            ]

        for i, (field, value) in enumerate(data_rows, row + 1):
            ws.cell(row=i, column=1, value=field)
            ws.cell(row=i, column=2, value=str(value) if value is not None else "N/A")

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 40

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        AuditService.log(
            db, action="report.download", user_id=user_id, ip_address=ip_address,
            resource_type="report", resource_id=prediction_id or project_id,
            description="Excel report generated", metadata={"format": "xlsx"},
        )
        return buffer.getvalue()

    @staticmethod
    def generate_pdf(
        db: Session,
        user_id: str,
        prediction_id: Optional[str] = None,
        project_id: Optional[str] = None,
        ip_address: Optional[str] = None,
    ) -> bytes:
        """Generate a PDF report with project details and SHAP explanation."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError:
            raise HTTPException(status_code=500, detail="reportlab is required for PDF reports.")

        prediction, project = ReportService._get_prediction_or_project(db, prediction_id, project_id)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph("RiskVision AI — Risk Assessment Report", styles["Title"]))
        elements.append(Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]))
        elements.append(Spacer(1, 20))

        if prediction:
            data = [
                ["Project ID", prediction.external_project_id],
                ["Project Name", prediction.project_name or "N/A"],
                ["Risk Level", prediction.risk_level],
                ["Risk Score", str(prediction.risk_score)],
                ["Failure Probability", f"{prediction.failure_probability:.1%}"],
                ["Prediction", prediction.prediction_label],
                ["Confidence", f"{prediction.confidence_level:.1%}"],
            ]
            table = Table(data, colWidths=[180, 300])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#1F4E79")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 20))

            if prediction.human_explanation:
                elements.append(Paragraph("Explanation", styles["Heading2"]))
                elements.append(Paragraph(prediction.human_explanation, styles["Normal"]))
                elements.append(Spacer(1, 12))

            if prediction.top_risk_factors:
                elements.append(Paragraph("Top Risk Factors", styles["Heading2"]))
                factor_data = [["Feature", "Impact", "Direction"]]
                for f in prediction.top_risk_factors[:5]:
                    if isinstance(f, dict):
                        factor_data.append([
                            f.get("display_name", f.get("feature_name", "")),
                            f"{f.get('impact', 0):.4f}",
                            f.get("direction", ""),
                        ])
                ft = Table(factor_data, colWidths=[200, 100, 180])
                ft.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
                elements.append(ft)

            if prediction.recommended_actions:
                elements.append(Spacer(1, 12))
                elements.append(Paragraph("Recommendations", styles["Heading2"]))
                for i, rec in enumerate(prediction.recommended_actions[:5], 1):
                    if isinstance(rec, dict):
                        elements.append(Paragraph(
                            f"{i}. [{rec.get('priority', '')}] {rec.get('action', '')}",
                            styles["Normal"],
                        ))

        doc.build(elements)
        buffer.seek(0)

        AuditService.log(
            db, action="report.download", user_id=user_id, ip_address=ip_address,
            resource_type="report", resource_id=prediction_id or project_id,
            description="PDF report generated", metadata={"format": "pdf"},
        )
        return buffer.getvalue()
